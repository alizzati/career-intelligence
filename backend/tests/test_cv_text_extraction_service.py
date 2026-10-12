from io import BytesIO
import struct
import zlib
from unittest import mock
import zipfile

import pymupdf
from docx import Document
import pytest

from app.agents.profile_agent import ProfileAgent
from app.services.cv_text_extraction_service import (
    CVTextExtractionError,
    CVTextExtractor,
    extract_cv_text,
)


def make_png() -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + kind
            + data
            + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
        )

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(b"\x00\xff\xff\xff\xff"))
        + chunk(b"IEND", b"")
    )


def make_pdf(*pages: str, encrypted: bool = False, image_only: bool = False) -> bytes:
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
        if image_only:
            page.insert_image(page.rect, stream=make_png())
    options = {}
    if encrypted:
        options = {
            "encryption": pymupdf.PDF_ENCRYPT_AES_256,
            "owner_pw": "owner-test-password",
            "user_pw": "user-test-password",
        }
    content = document.tobytes(**options)
    document.close()
    return content


def make_docx(*, with_table: bool = False, with_text: bool = True) -> bytes:
    document = Document()
    if with_text:
        document.add_paragraph("Synthetic candidate with backend development experience.")
    if with_table:
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Skills: Python and database management"
        document.sections[0].header.paragraphs[0].text = "Email: synthetic@example.test"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def error_code(call, expected_code: str) -> None:
    with pytest.raises(CVTextExtractionError) as exc_info:
        call()
    assert exc_info.value.code == expected_code


def test_extracts_synthetic_selectable_pdf_and_safe_metadata():
    content = make_pdf("Synthetic candidate profile with backend engineering experience.")

    result = extract_cv_text(content, "candidate.PDF")

    assert "Synthetic candidate profile" in result.extracted_text
    assert result.metadata.format == "pdf"
    assert result.metadata.size_bytes == len(content)
    assert result.metadata.page_count == 1
    assert result.metadata.character_count == len(result.extracted_text)
    assert not hasattr(result.metadata, "filename")


def test_extracts_docx_body_tables_and_header():
    content = make_docx(with_table=True)

    result = extract_cv_text(content, "candidate.docx")

    assert "Synthetic candidate" in result.extracted_text
    assert "Python and database management" in result.extracted_text
    assert "synthetic@example.test" in result.extracted_text
    assert result.metadata.format == "docx"
    assert result.metadata.page_count is None


@pytest.mark.parametrize(
    ("content", "expected"),
    [
        ("Plain UTF-8 CV text with synthetic content.".encode(), "Plain UTF-8"),
        (b"\xef\xbb\xbfUTF-8 BOM candidate profile.", "UTF-8 BOM"),
        ("Synthetic UTF-16 profile".encode("utf-16"), "Synthetic UTF-16"),
    ],
)
def test_extracts_safely_decoded_txt(content, expected):
    result = extract_cv_text(content, "candidate.txt")

    assert expected in result.extracted_text
    assert result.metadata.format == "txt"
    assert result.metadata.page_count is None


def test_rejects_empty_upload_and_empty_supported_documents():
    error_code(lambda: extract_cv_text(b"", "empty.txt"), "empty_file")
    error_code(lambda: extract_cv_text(b" \r\n\t", "empty.txt"), "empty_document")
    error_code(
        lambda: extract_cv_text(make_docx(with_text=False), "empty.docx"),
        "empty_document",
    )


def test_rejects_malformed_pdf_docx_and_unsafe_text_encoding():
    error_code(lambda: extract_cv_text(b"%PDF-1.7 broken", "broken.pdf"), "unreadable_document")
    error_code(lambda: extract_cv_text(b"PK\x03\x04broken", "broken.docx"), "unreadable_document")
    error_code(lambda: extract_cv_text(b"\xff\xfe\x00", "broken.txt"), "unsupported_text_encoding")
    error_code(lambda: extract_cv_text(b"binary\x00payload", "binary.txt"), "unreadable_document")
    error_code(lambda: extract_cv_text(b"binary\x1bpayload", "binary.txt"), "unreadable_document")


@pytest.mark.parametrize(
    ("content", "filename"),
    [
        (b"%PDF-1.7", "disguised.txt"),
        (b"synthetic text", "fake.pdf"),
        (b"PK\x03\x04not-docx", "fake.docx"),
        (make_docx(), "misnamed.pdf"),
        (b"plain synthetic text", "unsupported.docm"),
    ],
    ids=["pdf-as-txt", "text-as-pdf", "malformed-docx", "docx-as-pdf", "unsupported-extension"],
)
def test_rejects_incorrect_extensions_or_signatures(content, filename):
    expected = {
        "unsupported.docm": "unsupported_format",
        "fake.docx": "unreadable_document",
    }.get(filename, "format_mismatch")
    error_code(lambda: extract_cv_text(content, filename), expected)


def test_enforces_configurable_file_size_and_pdf_page_limits():
    size_limited = CVTextExtractor(max_file_size_bytes=10)
    error_code(
        lambda: size_limited.extract(b"larger than limit", "candidate.txt"),
        "file_too_large",
    )

    page_limited = CVTextExtractor(max_pdf_pages=1)
    error_code(
        lambda: page_limited.extract(
            make_pdf(
                "First page has meaningful synthetic content.",
                "Second page has meaningful synthetic content.",
            ),
            "candidate.pdf",
        ),
        "page_limit_exceeded",
    )


def test_rejects_encrypted_pdf_and_image_only_scanned_pdf():
    error_code(
        lambda: extract_cv_text(make_pdf("Synthetic encrypted content", encrypted=True), "locked.pdf"),
        "encrypted_pdf",
    )
    with pytest.raises(CVTextExtractionError) as exc_info:
        extract_cv_text(make_pdf("", image_only=True), "scanned.pdf")
    assert exc_info.value.code == "scanned_pdf_unsupported"
    assert "Scanned/image-only PDFs are unsupported" in str(exc_info.value)


def test_rejects_docx_with_excessive_uncompressed_size():
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"x" * 256)
        archive.writestr("word/document.xml", b"<document/>")
    extractor = CVTextExtractor(max_docx_uncompressed_bytes=100)

    error_code(
        lambda: extractor.extract(buffer.getvalue(), "expanded.docx"),
        "document_expansion_limit_exceeded",
    )


def test_extraction_does_not_call_profile_agent_write_files_or_log_document_text(
    tmp_path,
    caplog,
):
    content_text = "Synthetic private@example.test profile text; never log this document."
    content = content_text.encode("utf-8")
    with mock.patch.object(ProfileAgent, "run") as profile_run:
        result = extract_cv_text(content, "sensitive-looking-filename.txt")

    assert result.extracted_text == content_text
    profile_run.assert_not_called()
    assert list(tmp_path.iterdir()) == []
    assert content_text not in caplog.text
    assert "sensitive-looking-filename.txt" not in caplog.text


def test_extraction_errors_are_sanitized_and_do_not_echo_document_or_filename(caplog):
    sensitive_filename = "private@example.test.pdf"
    sensitive_document = b"%PDF-1.7 PRIVATE-IDENTIFIER"
    with pytest.raises(CVTextExtractionError) as exc_info:
        extract_cv_text(sensitive_document, sensitive_filename)

    assert "PRIVATE-IDENTIFIER" not in str(exc_info.value)
    assert "private@example.test" not in str(exc_info.value)
    assert "PRIVATE-IDENTIFIER" not in caplog.text
    assert "private@example.test" not in caplog.text


@pytest.mark.parametrize(
    ("kwargs", "field"),
    [
        ({"max_file_size_bytes": 0}, "max_file_size_bytes"),
        ({"max_pdf_pages": True}, "max_pdf_pages"),
        ({"max_docx_uncompressed_bytes": -1}, "max_docx_uncompressed_bytes"),
    ],
)
def test_rejects_invalid_resource_limits(kwargs, field):
    with pytest.raises(ValueError, match=field):
        CVTextExtractor(**kwargs)
