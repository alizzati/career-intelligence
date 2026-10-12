"""In-memory text extraction for supported CV document formats.

Extracted text is returned only to the caller. This service does not persist
documents or text, log content, call an LLM, or apply PII redaction.
"""
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath
import re
import unicodedata
import zipfile

import pymupdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


DEFAULT_MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024
DEFAULT_MAX_PDF_PAGES = 20
DEFAULT_MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


class CVTextExtractionError(ValueError):
    """Sanitized document validation or extraction failure."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class CVDocumentMetadata:
    format: str
    size_bytes: int
    page_count: int | None
    character_count: int


@dataclass(frozen=True)
class CVTextExtraction:
    extracted_text: str
    metadata: CVDocumentMetadata


class CVTextExtractor:
    """Extract PDF, DOCX, and safely decoded TXT content from in-memory bytes."""

    def __init__(
        self,
        *,
        max_file_size_bytes: int = DEFAULT_MAX_FILE_SIZE_BYTES,
        max_pdf_pages: int = DEFAULT_MAX_PDF_PAGES,
        max_docx_uncompressed_bytes: int = DEFAULT_MAX_DOCX_UNCOMPRESSED_BYTES,
    ):
        for name, value in (
            ("max_file_size_bytes", max_file_size_bytes),
            ("max_pdf_pages", max_pdf_pages),
            ("max_docx_uncompressed_bytes", max_docx_uncompressed_bytes),
        ):
            if not isinstance(value, int) or isinstance(value, bool) or value <= 0:
                raise ValueError(f"{name} must be a positive integer")
        self.max_file_size_bytes = max_file_size_bytes
        self.max_pdf_pages = max_pdf_pages
        self.max_docx_uncompressed_bytes = max_docx_uncompressed_bytes

    def extract(self, file_bytes: bytes, filename: str) -> CVTextExtraction:
        """Validate metadata/signature and extract text without storing input."""
        if not isinstance(file_bytes, bytes):
            raise CVTextExtractionError("invalid_input", "Document content must be bytes.")
        if not file_bytes:
            raise CVTextExtractionError("empty_file", "The uploaded document is empty.")
        if len(file_bytes) > self.max_file_size_bytes:
            raise CVTextExtractionError(
                "file_too_large",
                f"Document exceeds the {self.max_file_size_bytes}-byte file-size limit.",
            )
        if not isinstance(filename, str) or not filename.strip():
            raise CVTextExtractionError("invalid_filename", "A supported filename is required.")

        extension = PurePath(filename).suffix.lower()
        if extension not in {".pdf", ".docx", ".txt"}:
            raise CVTextExtractionError(
                "unsupported_format",
                "Unsupported document format; provide a PDF, DOCX, or TXT file.",
            )

        self._validate_signature(extension, file_bytes)
        if extension == ".pdf":
            extracted_text, page_count = self._extract_pdf(file_bytes)
        elif extension == ".docx":
            extracted_text = self._extract_docx(file_bytes)
            page_count = None
        else:
            extracted_text = self._decode_text(file_bytes)
            page_count = None

        if not extracted_text.strip():
            raise CVTextExtractionError(
                "empty_document",
                "No text was found in the document.",
            )

        return CVTextExtraction(
            extracted_text=extracted_text,
            metadata=CVDocumentMetadata(
                format=extension[1:],
                size_bytes=len(file_bytes),
                page_count=page_count,
                character_count=len(extracted_text),
            ),
        )

    @staticmethod
    def _validate_signature(extension: str, file_bytes: bytes) -> None:
        has_pdf_header = b"%PDF-" in file_bytes[:1024]
        if extension == ".pdf" and not has_pdf_header:
            raise CVTextExtractionError(
                "format_mismatch",
                "The PDF filename does not match the document content.",
            )
        if extension == ".docx":
            if not file_bytes.startswith(b"PK\x03\x04"):
                raise CVTextExtractionError(
                    "format_mismatch",
                    "The DOCX filename does not match the document content.",
                )
            try:
                with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
                    names = set(archive.namelist())
                    if (
                        "[Content_Types].xml" not in names
                        or "word/document.xml" not in names
                    ):
                        raise CVTextExtractionError(
                            "format_mismatch",
                            "The ZIP document is not a valid DOCX file.",
                        )
                    if any(name.lower().endswith("vbaproject.bin") for name in names):
                        raise CVTextExtractionError(
                            "unsupported_format",
                            "Macro-enabled documents are not supported.",
                        )
            except CVTextExtractionError:
                raise
            except Exception:
                raise CVTextExtractionError(
                    "unreadable_document",
                    "The DOCX file is malformed or unreadable.",
                ) from None
        if extension == ".txt" and (
            has_pdf_header or file_bytes.startswith(b"PK\x03\x04")
        ):
            raise CVTextExtractionError(
                "format_mismatch",
                "The TXT filename does not match the document content.",
            )

    def _extract_pdf(self, file_bytes: bytes) -> tuple[str, int]:
        try:
            with pymupdf.open(stream=file_bytes, filetype="pdf") as document:
                if document.needs_pass:
                    raise CVTextExtractionError(
                        "encrypted_pdf",
                        "Password-protected PDFs are not supported; provide an unlocked PDF.",
                    )
                page_count = document.page_count
                if page_count > self.max_pdf_pages:
                    raise CVTextExtractionError(
                        "page_limit_exceeded",
                        f"PDF exceeds the {self.max_pdf_pages}-page limit.",
                    )
                text = "\n".join(page.get_text("text") for page in document)
        except CVTextExtractionError:
            raise
        except Exception:
            raise CVTextExtractionError(
                "unreadable_document",
                "The PDF is malformed or unreadable.",
            ) from None

        alphanumeric_count = sum(char.isalnum() for char in text)
        words = re.findall(r"[A-Za-z0-9]{2,}", text)
        if alphanumeric_count < 10 or len(words) < 2:
            raise CVTextExtractionError(
                "scanned_pdf_unsupported",
                "No meaningful selectable text was found. Scanned/image-only PDFs are "
                "unsupported; provide a text-based PDF or DOCX/TXT file.",
            )
        return text, page_count

    def _extract_docx(self, file_bytes: bytes) -> str:
        try:
            with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
                uncompressed_size = sum(item.file_size for item in archive.infolist())
                if uncompressed_size > self.max_docx_uncompressed_bytes:
                    raise CVTextExtractionError(
                        "document_expansion_limit_exceeded",
                        "DOCX expanded content exceeds the configured safety limit.",
                    )

            document = Document(BytesIO(file_bytes))
            chunks: list[str] = []
            for child in document.element.body.iterchildren():
                if child.tag.endswith("}p"):
                    chunks.append(Paragraph(child, document).text)
                elif child.tag.endswith("}tbl"):
                    self._append_table_text(Table(child, document), chunks)

            for section in document.sections:
                for region in (section.header, section.footer):
                    chunks.extend(paragraph.text for paragraph in region.paragraphs)
                    for table in region.tables:
                        self._append_table_text(table, chunks)
            return "\n".join(part for part in chunks if part)
        except CVTextExtractionError:
            raise
        except Exception:
            raise CVTextExtractionError(
                "unreadable_document",
                "The DOCX file is malformed or unreadable.",
            ) from None

    @classmethod
    def _append_table_text(cls, table: Table, chunks: list[str]) -> None:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(paragraph.text for paragraph in cell.paragraphs)
                for nested_table in cell.tables:
                    cls._append_table_text(nested_table, chunks)

    @staticmethod
    def _decode_text(file_bytes: bytes) -> str:
        try:
            if file_bytes.startswith((b"\xff\xfe", b"\xfe\xff")):
                text = file_bytes.decode("utf-16")
            else:
                text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise CVTextExtractionError(
                "unsupported_text_encoding",
                "TXT must use valid UTF-8 or BOM-marked UTF-16 encoding.",
            ) from None
        if any(
            unicodedata.category(char) == "Cc" and char not in "\t\n\r\f"
            for char in text
        ):
            raise CVTextExtractionError(
                "unreadable_document",
                "The TXT file contains binary data and cannot be decoded safely.",
            )
        return text


def extract_cv_text(
    file_bytes: bytes,
    filename: str,
    *,
    max_file_size_bytes: int = DEFAULT_MAX_FILE_SIZE_BYTES,
    max_pdf_pages: int = DEFAULT_MAX_PDF_PAGES,
) -> CVTextExtraction:
    """Convenience interface using the default DOCX expansion safety limit."""
    return CVTextExtractor(
        max_file_size_bytes=max_file_size_bytes,
        max_pdf_pages=max_pdf_pages,
    ).extract(file_bytes, filename)
